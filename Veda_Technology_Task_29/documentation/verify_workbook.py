from pathlib import Path
import csv
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
with (root / "raw_data" / "retail_sales.csv").open(encoding="utf-8", newline="") as f:
    data = list(csv.DictReader(f))
wb = load_workbook(root / "workbook" / "Retail_Sales_Conditional_Formatting.xlsx")
ws = wb["Retail Sales"]
assert len(data) == ws.max_row - 1
assert "Rules List" in wb.sheetnames
assert len(ws.conditional_formatting) >= 4
print(f"Verification passed: {len(data)} data rows; sheets: {wb.sheetnames}")
print(f"Conditional formatting ranges: {list(ws.conditional_formatting)}")
